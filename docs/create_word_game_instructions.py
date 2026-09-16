from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUTPUT = "output/How to Play The Word Discovery Challenge.docx"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color="D9D9D9", size="8"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:color"), color)


def set_cell_margins(cell, top=110, start=140, bottom=110, end=140):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn("w:" + margin))
        if node is None:
            node = OxmlElement("w:" + margin)
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def mark_header_row(row):
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def keep_with_next(paragraph):
    p_pr = paragraph._p.get_or_add_pPr()
    keep = p_pr.find(qn("w:keepNext"))
    if keep is None:
        keep = OxmlElement("w:keepNext")
        p_pr.append(keep)


def keep_lines(paragraph):
    p_pr = paragraph._p.get_or_add_pPr()
    keep = p_pr.find(qn("w:keepLines"))
    if keep is None:
        keep = OxmlElement("w:keepLines")
        p_pr.append(keep)


def set_font(run, name="Liberation Sans"):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)


def remove_paragraph_border(paragraph):
    p_pr = paragraph._p.get_or_add_pPr()
    border = p_pr.find(qn("w:pBdr"))
    if border is not None:
        p_pr.remove(border)


def remove_style_border(style):
    p_pr = style._element.get_or_add_pPr()
    border = p_pr.find(qn("w:pBdr"))
    if border is not None:
        p_pr.remove(border)


def add_step(doc, number, heading, body):
    paragraph = doc.add_paragraph(style="List Number")
    keep_lines(paragraph)
    paragraph.paragraph_format.space_after = Pt(5)
    paragraph.paragraph_format.left_indent = Inches(0.24)
    paragraph.paragraph_format.first_line_indent = Inches(-0.24)
    title = paragraph.add_run(heading + "  ")
    title.bold = True
    title.font.color.rgb = RGBColor(0, 0, 0)
    set_font(title)
    text = paragraph.add_run(body)
    set_font(text)
    return paragraph


def add_bullet(doc, lead, body):
    paragraph = doc.add_paragraph(style="List Bullet")
    keep_lines(paragraph)
    paragraph.paragraph_format.space_after = Pt(3)
    paragraph.paragraph_format.left_indent = Inches(0.24)
    paragraph.paragraph_format.first_line_indent = Inches(-0.18)
    run = paragraph.add_run(lead)
    run.bold = True
    set_font(run)
    run = paragraph.add_run(body)
    set_font(run)
    return paragraph


doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(0.62)
section.bottom_margin = Inches(0.58)
section.left_margin = Inches(0.78)
section.right_margin = Inches(0.78)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Liberation Sans"
normal._element.rPr.rFonts.set(qn("w:ascii"), "Liberation Sans")
normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Liberation Sans")
normal.font.size = Pt(10.8)
normal.font.color.rgb = RGBColor(28, 28, 28)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.05

for style_name, size, before, after in (
    ("Title", 26, 0, 5),
    ("Subtitle", 11, 0, 14),
    ("Heading 1", 15, 11, 5),
    ("Heading 2", 11.5, 7, 3),
):
    style = styles[style_name]
    style.font.name = "Liberation Sans"
    style._element.rPr.rFonts.set(qn("w:ascii"), style.font.name)
    style._element.rPr.rFonts.set(qn("w:hAnsi"), style.font.name)
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.font.bold = style_name != "Subtitle"
    style.paragraph_format.space_before = Pt(before)
    style.paragraph_format.space_after = Pt(after)
    remove_style_border(style)

title = doc.add_paragraph("How to Play The Word Discovery Challenge", style="Title")
title.alignment = WD_ALIGN_PARAGRAPH.LEFT
remove_paragraph_border(title)
subtitle = doc.add_paragraph("A concise guide to guesses, confidence ratings, clues, timing, and results", style="Subtitle")
remove_paragraph_border(subtitle)
subtitle.runs[0].font.color.rgb = RGBColor(78, 78, 78)

intro = doc.add_paragraph()
lead = intro.add_run("Goal  ")
lead.bold = True
set_font(lead)
run = intro.add_run("Find each hidden five-letter word. You will play one practice word, then 20 scored words.")
set_font(run)

timing = doc.add_paragraph()
lead = timing.add_run("Time limit  ")
lead.bold = True
set_font(lead)
run = timing.add_run("You have 100 minutes for the practice word and all 20 scored words together. The timer does not pause between words or when you leave the page.")
set_font(run)

heading = doc.add_paragraph("Step by step", style="Heading 1")
keep_with_next(heading)

add_step(doc, 1, "Start with the practice word.", "It works like every later word, but it does not count toward your score. The shared timer starts with practice.")
add_step(doc, 2, "Enter a five-letter word.", "Type with your keyboard or use the on-screen keys, then select Submit guess. A word that is not in the game dictionary does not use one of your attempts.")
add_step(doc, 3, "Rate your confidence before seeing the clues.", "Choose 0 to 10 for the chance that this exact guess is the hidden word. For example, 0 means 0%, 5 means 50%, and 10 means 100%.")
add_step(doc, 4, "Confirm the rating.", "Select Confirm confidence and show clues. You cannot enter another guess until you confirm the rating.")
add_step(doc, 5, "Use the color clues.", "Compare the letters in your guess with the hidden word, using the key below.")

table = doc.add_table(rows=1, cols=3)
table.autofit = False
mark_header_row(table.rows[0])
table.columns[0].width = Inches(1.15)
table.columns[1].width = Inches(1.55)
table.columns[2].width = Inches(4.2)
headers = ("Color", "Meaning", "What to do")
for index, text in enumerate(headers):
    cell = table.rows[0].cells[index]
    set_cell_shading(cell, "1F2937")
    set_cell_border(cell)
    set_cell_margins(cell)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    paragraph = cell.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if index < 2 else WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.add_run(text)
    run.bold = True
    run.font.color.rgb = RGBColor(255, 255, 255)
    run.font.size = Pt(10)
    set_font(run)

clues = [
    ("Green", "Correct position", "Keep the letter in the same position in a later guess."),
    ("Yellow", "Different position", "The letter is in the word; try it in another position."),
    ("Gray", "Not matched", "The letter is not matched in the hidden word for that guess."),
]
fills = {"Green": "6AAA64", "Yellow": "C9B458", "Gray": "787C7E"}
for row_index, (color, meaning, action) in enumerate(clues, start=1):
    cells = table.add_row().cells
    values = (color, meaning, action)
    for col_index, value in enumerate(values):
        cell = cells[col_index]
        set_cell_border(cell)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        if col_index == 0:
            set_cell_shading(cell, fills[color])
        else:
            set_cell_shading(cell, "F7F8FA" if row_index % 2 else "FFFFFF")
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if col_index < 2 else WD_ALIGN_PARAGRAPH.LEFT
        paragraph.paragraph_format.space_after = Pt(0)
        run = paragraph.add_run(value)
        run.font.size = Pt(9.8)
        run.bold = col_index == 0
        run.font.color.rgb = RGBColor(255, 255, 255) if col_index == 0 else RGBColor(20, 20, 20)
        set_font(run)

after_table = doc.add_paragraph()
after_table.paragraph_format.space_before = Pt(7)
after_table.paragraph_format.space_after = Pt(5)
run = after_table.add_run("Note  ")
run.bold = True
set_font(run)
run = after_table.add_run("Each new guess is entered from scratch. Green letters are not automatically copied or locked, so reuse them in their correct positions yourself.")
set_font(run)

add_step(doc, 6, "Keep guessing.", "Repeat the guess, confidence, and clue steps until you solve the word or use all 10 valid guesses.")
add_step(doc, 7, "Move to the next word.", "After a round ends, the answer is shown. Select Next word. After practice, select Finish practice and then Continue to begin the 20 scored words.")
add_step(doc, 8, "Finish and save your results.", "After the twentieth word, select See results. Download the JSON or CSV file before selecting Play again, because results are stored only in this browser.")

heading = doc.add_paragraph("Important rules", style="Heading 1")
keep_with_next(heading)
add_bullet(doc, "Ten attempts maximum.  ", "Only valid dictionary words use an attempt.")
add_bullet(doc, "Confidence comes first.  ", "The clues remain hidden until you confirm a confidence rating for the guess.")
add_bullet(doc, "Giving up.  ", "You may select Give up this word when no confidence rating is waiting for confirmation. The round counts as unsuccessful.")
add_bullet(doc, "When time expires.  ", "The game ends automatically. Unconfirmed guesses are not counted, and words you did not reach are recorded separately.")
add_bullet(doc, "Language entry.  ", "Use A-Z in English. In German, enter Ä as AE, Ö as OE, Ü as UE, and ß as SS; the converted word must fill exactly five spaces.")

heading = doc.add_paragraph("Quick play loop", style="Heading 1")
keep_with_next(heading)
loop = doc.add_paragraph()
loop.alignment = WD_ALIGN_PARAGRAPH.CENTER
loop.paragraph_format.space_before = Pt(1)
loop.paragraph_format.space_after = Pt(0)
run = loop.add_run("GUESS  →  RATE CONFIDENCE  →  CONFIRM  →  READ CLUES  →  REPEAT")
run.bold = True
run.font.size = Pt(10.5)
run.font.color.rgb = RGBColor(31, 41, 55)
set_font(run)

footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
footer.paragraph_format.space_before = Pt(0)
run = footer.add_run("The Word Discovery Challenge quick guide")
run.font.size = Pt(8.5)
run.font.color.rgb = RGBColor(100, 100, 100)
set_font(run)

doc.core_properties.title = "How to Play The Word Discovery Challenge"
doc.core_properties.subject = "Step-by-step instructions for The Word Discovery Challenge"
doc.core_properties.author = ""
doc.core_properties.keywords = "The Word Discovery Challenge, instructions, confidence, clues"
doc.save(OUTPUT)
print(OUTPUT)
